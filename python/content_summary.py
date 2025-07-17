import json
import logging
import warnings
import constants
import configparser
from utils import timed
from methods import line_number
from constants import UTF8, READ, HASH_MD5, MAIN, SETTINGS_FILE

warnings.filterwarnings('ignore')
logging.getLogger('transformers').setLevel(logging.ERROR)
show = logging.getLogger(constants.DEBUG_EVAL)

OLLAMA_MODEL = 'tinyllama'  # phi3 # phi3:mini

processor = None
model = None

config = configparser.ConfigParser()
config.read(SETTINGS_FILE)

openai_key = config.get('SECURITY', 'openai')

TEXT_2_SUMMARY = 'TAG__TXT_2_SUMMARY'
AUDIO_2_TXT = 'TAG__AUDIO_2_TXT'
IMAGE_2_DESCRIPTION = 'TAG__IMAGE_2_DESCRIPTION'
OCR_2_TXT = 'TAG__OCR_2_TXT'
VIDEO_2_DESCRIPTION = 'TAG__VIDEO_2_DESCRIPTION'
VIDEO_2_TXT = 'TAG__VIDEO_2_TXT'
LLAMA = 'TAG__LLAMA'
OPENAI = 'TAG__OPENAI'
RAW_TEXT = 'TAG__RAW_TEXT'
PDF_2_TXT = 'TAG__PDF_2_TXT'
DOCX_TO_TXT = 'TAG__DOCX_2_TXT'
XLSX_2_TXT = 'TAG__XLSX_2_TXT'
PPTX_2_TXT = 'TAG__PPTX_2_TXT'
RTF_2_TXT = 'TAG__RTF_2_TXT'
XLS_2_TXT = 'TAG__XLS_2_TXT'
HTML_2_TXT = 'TAG__HTML_2_TXT'
IPYNB_2_TXT = 'TAG__IPYNB_2_TXT'
DOC_2_TXT = 'TAG__DOC_2_TXT'
ZIP_2_TXT = 'TAG__ZIP_2_TXT'
EMAIL_SUBJECT = 'TAG__SUBJECT'
EMAIL_BODY = 'TAG__BODY'
DBX_2_TXT = 'TAG__DBX_2_TXT'
MARKIT = 'TAG__MARKIT'

def normalize(text):
    return text.replace('\n', ' ').replace('\t', ' ').strip()

@timed
def txt_to_summary(text_to_summarize):
    show.error(f'{line_number()} [Summarizing]...')
    import requests
    try:
        answer_timeout = 240
        response = requests.post('http://localhost:11434/api/generate', json={'model': OLLAMA_MODEL, 'prompt': f'Summarize the following:\n\n{text_to_summarize}', 'stream': False}, timeout=answer_timeout)
        return {TEXT_2_SUMMARY: response.json().get('response', '')}
    except Exception:
        return {TEXT_2_SUMMARY: None}

def mark(path):
    if path.endswith(('.csv', '.docx', '.epub', '.htm', '.html', '.jpg', '.jpeg', '.png', '.bmp', '.webp', '.json', '.mht', '.mhtml', '.mp3', '.outlook', '.pdf', '.pptx', '.tsv', '.txt', '.wav', '.xls', '.xlsx', '.xml', '.zip')):
        result = mark_it(path)
        if result['MARKIT[OFF]'] and result['MARKIT[OFF]'] != '':
            return result
        else:
            return mark_it_llama(path)
    else:
        return {MARKIT: None}

def mark_it(path, plugins=False):
    show.error(f'{line_number()} [Plugins Marking]...')
    from markitdown import MarkItDown
    status = 'ON' if plugins else 'OFF'
    try:
        md = MarkItDown(enable_plugins=plugins)
        result = md.convert(path)
        return {f'MARKIT[{status}]': normalize(result.text_content)}
    except Exception as exception:
        show.error(f'{line_number()} ERROR [MarkIt Plugins] [{exception}]')
        return {f'MARKIT[{status}]': None}

def mark_it_llama(path):
    show.error(f'{line_number()} [Llama Marking]...')
    from markitdown import MarkItDown
    try:
        md = MarkItDown(docintel_endpoint='http://localhost:11434/api/generate')
        result = md.convert(path)
        return {LLAMA: result.text_content}
    except Exception as exception:
        show.error(f'{line_number()} ERROR [MarkIt Llama] [{exception}]')
        return {LLAMA: None}

def mark_it_openai(path):
    show.error(f'{line_number()} [OpenAI Marking]...')
    from openai import OpenAI
    from markitdown import MarkItDown
    try:
        client = OpenAI(api_key=openai_key)
        md = MarkItDown(llm_client=client, llm_model='gpt-3.5-turbo')
        result = md.convert(path)
        return {OPENAI: result.text_content}
    except Exception as exception:
        show.error(f'{line_number()} ERROR [MarkIt OpenAI] [{exception}]')
        return {OPENAI: None}

@timed
def audio_to_txt(path):
    import whisper
    try:
        local_model = whisper.load_model('base')
        return {AUDIO_2_TXT: local_model.transcribe(path)['text']}
    except:
        return {AUDIO_2_TXT: None}

def image_to_description(image_path):
    show.error(f'{line_number()} [Describing {image_path}]...')
    global model
    global processor
    if not processor:
        from transformers import BlipProcessor, BlipForConditionalGeneration
        processor = BlipProcessor.from_pretrained('Salesforce/blip-image-captioning-base')
        model = BlipForConditionalGeneration.from_pretrained('Salesforce/blip-image-captioning-base')
    from PIL import Image
    image = Image.open(image_path).convert('RGB')
    inputs = processor(image, return_tensors='pt')
    out = model.generate(**inputs)
    return {IMAGE_2_DESCRIPTION: processor.decode(out[0], skip_special_tokens=True)}

def ocr_to_txt(path):
    import pytesseract
    from PIL import Image
    return {OCR_2_TXT: normalize(pytesseract.image_to_string(Image.open(path)))}

def image_to_txt(path):
    return ocr_to_txt(path) | image_to_description(path)

@timed
def video_to_txt(path):
    import os
    import cv2
    temp_dir = 'frames'
    os.makedirs(temp_dir, exist_ok=True)
    cap = cv2.VideoCapture(path)
    count, frame_count = 0, int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    descriptions = []
    while count < frame_count:
        cap.set(cv2.CAP_PROP_POS_FRAMES, count)
        ret, frame = cap.read()
        if not ret:
            break
        frame_path = os.path.join(temp_dir, f'frame_{count}.jpg')
        cv2.imwrite(frame_path, frame)
        descriptions.append(ocr_to_txt(frame_path)[OCR_2_TXT])
        count += int(frame_count / 5)
    cap.release()
    return {VIDEO_2_TXT: normalize('\n'.join(descriptions))}

@timed
def video_to_description(path):
    import os
    import cv2
    temp_dir = 'frames'
    os.makedirs(temp_dir, exist_ok=True)
    cap = cv2.VideoCapture(path)
    count, frame_count = 0, int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    descriptions = []
    while count < frame_count:
        cap.set(cv2.CAP_PROP_POS_FRAMES, count)
        ret, frame = cap.read()
        if not ret:
            break
        frame_path = os.path.join(temp_dir, f'frame_{count}.jpg')
        cv2.imwrite(frame_path, frame)
        descriptions.append(image_to_description(frame_path)[IMAGE_2_DESCRIPTION])
        count += int(frame_count / 5)
    cap.release()
    return {VIDEO_2_DESCRIPTION: normalize('\n'.join(descriptions))}

def audio_video_to_txt(path):
    return audio_to_txt(path) | video_to_description(path) | video_to_txt(path)

def exe_to_txt(path):
    import os
    import pefile
    info = {}
    info_caps = {}
    os_info = os.stat(path)
    try:
        pe = pefile.PE(path)
        if hasattr(pe, 'FileInfo'):
            for fileinfo_list in pe.FileInfo:
                for fileinfo in fileinfo_list:
                    if hasattr(fileinfo, 'Key') and fileinfo.Key == b'StringFileInfo':
                        for st in fileinfo.StringTable:
                            for key, value in st.entries.items():
                                info[key.decode()] = value.decode(errors='ignore')
    except pefile.PEFormatError:
        info['error'] = 'Not a valid PE file'
    info['size'] = os_info.st_size
    for k in info:
        info_caps[k.upper()] = info[k]
    return info_caps

def read_txt(path):
    with open(path, READ, encoding=UTF8, errors='replace') as file:
        return {RAW_TEXT: normalize(file.read())}

def pdf_to_txt(path):
    from pdfminer.high_level import extract_text as extract_pdf
    return {PDF_2_TXT: normalize(extract_pdf(path))}

def docx_to_txt(path):
    from docx import Document
    try:
        doc = Document(path)
        return {DOCX_TO_TXT: normalize('\n'.join(p.text for p in doc.paragraphs))}
    except Exception:
        return {DOCX_TO_TXT: None}

def xlsx_to_txt(path):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, data_only=True)
        return {XLSX_2_TXT: normalize('\n'.join(str(cell.value) for ws in wb for row in ws.iter_rows() for cell in row if cell.value))}
    except Exception:
        return {XLSX_2_TXT: None}

def pptx_to_txt(path):
    import pptx
    try:
        prs = pptx.Presentation(path)
        text_runs = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    text_runs.append(shape.text)
        return {PPTX_2_TXT: normalize('\n'.join(text_runs))}
    except Exception:
        return {PPTX_2_TXT: None}

def rtf_to_txt(path):
    from striprtf.striprtf import rtf_to_text
    try:
        with open(path, READ, encoding=UTF8, errors='replace') as f:
            return {RTF_2_TXT: normalize(rtf_to_text(f.read()))}
    except Exception as exception:
        show.error(f'{line_number()} [{exception}]...')
        return {RTF_2_TXT: None}

def xls_to_txt(path):
    import xlrd
    try:
        wb = xlrd.open_workbook(path)
        result = '\n'.join(str(wb.sheet_by_index(0).cell_value(row, col)) for row in range(wb.sheet_by_index(0).nrows) for col in range(wb.sheet_by_index(0).ncols))
        return {XLS_2_TXT: normalize(result)}
    except:
        return {XLS_2_TXT: None}

def html_to_txt(path):
    from bs4 import BeautifulSoup
    with open(path, READ, encoding=UTF8, errors='replace') as f:
        soup = BeautifulSoup(f, 'html.parser')
        return {HTML_2_TXT: normalize(soup.get_text())}

def ipynb_to_txt(path):
    with open(path, READ, encoding=UTF8, errors='replace') as f:
        notebook = json.load(f)
        result = '\n'.join(cell['source'] if isinstance(cell['source'], str) else ''.join(cell['source']) for cell in notebook.get('cells', []) if cell.get('cell_type') in ('markdown', 'code'))
        return {IPYNB_2_TXT: normalize(result)}

def doc_to_txt(path):
    import textract
    try:
        return {DOC_2_TXT: normalize(textract.process(path).decode(UTF8, errors='replace'))}
    except Exception as exception:
        show.error(f'{line_number()} [{exception}]...')
        return {DOC_2_TXT: None}


def apk_to_txt(path):
    import os
    import zipfile
    from xml.dom import minidom
    info = {}
    os_info = os.stat(path)
    try:
        with zipfile.ZipFile(path, 'r') as zip_ref:
            if 'AndroidManifest.xml' in zip_ref.namelist():
                manifest = zip_ref.read('AndroidManifest.xml')
                try:
                    xml = minidom.parseString(manifest)
                except Exception:
                    info['ERROR'] = 'Invalid XML format'
                    xml = None
                if xml:
                    manifest_tag = xml.getElementsByTagName('manifest')[0]
                    info['package'] = manifest_tag.getAttribute('package')
                    if manifest_tag.hasAttribute('android:versionName'):
                        info['versionName'] = manifest_tag.getAttribute('android:versionName')
                    if manifest_tag.hasAttribute('android:versionCode'):
                        info['versionCode'] = manifest_tag.getAttribute('android:versionCode')
    except zipfile.BadZipFile:
        info['ERROR'] = 'Not a valid APK file'
    info['SIZE'] = os_info.st_size
    return info

def msg_to_txt(path):
    import pythoncom
    import win32com.client
    try:
        pythoncom.CoInitialize()
        outlook = win32com.client.Dispatch('Outlook.Application').GetNamespace('MAPI')
        msg = outlook.OpenSharedItem(path)
        msg_body = normalize(str(msg.Body))
        return {EMAIL_SUBJECT: msg.Subject, EMAIL_BODY: msg_body}
    except:
        return {EMAIL_SUBJECT: None, EMAIL_BODY: None}

def dbx_to_txt(path):
    import shutil
    import mailbox
    try:
        temp_mbox = 'temp.mbox'
        shutil.copy(path, temp_mbox)
        mbox = mailbox.mbox(temp_mbox)
        messages = []
        for msg in mbox:
            messages.append(msg.get('subject', '') + '\n' + msg.get_payload(decode=True).decode(errors='replace'))
        return {DBX_2_TXT: '\n\n'.join(messages)}
    except:
        return {DBX_2_TXT: None}

def eml_to_txt(path):
    from email import policy
    from email.parser import BytesParser
    try:
        with open(path, 'rb') as f:
            msg = BytesParser(policy=policy.default).parse(f)
        return {EMAIL_SUBJECT: msg['subject'], EMAIL_BODY: msg.get_body(preferencelist="plain").get_content()}
    except:
        return None

def zip_to_txt(path):
    import os
    import zipfile
    with zipfile.ZipFile(path, 'r') as z:
        result = {}
        for name in z.namelist():
            if any(name.lower().endswith(ext) for t in TYPES for ext in t if ext != 'category'):
                z.extract(name, 'temp')
                result[name] = run_function(os.path.join('temp', name))
    return {ZIP_2_TXT: result}

TYPES = [
    {'.7z': None, 'category': 'compressed'},
    {'.aac': audio_to_txt, 'category': 'audio'},
    {'.apk': apk_to_txt, 'category': 'android'},
    {'.arj': None, 'category': 'compressed'},
    {'.au': audio_to_txt, 'category': 'audio'},
    {'.avi': audio_video_to_txt, 'category': 'video'},
    {'.avif': image_to_txt, 'category': 'image'},
    {'.bak': read_txt, 'category': 'text'},
    {'.bat': read_txt, 'category': 'text'},
    {'.bmp': image_to_txt, 'category': 'image'},
    {'.c': read_txt, 'category': 'text'},
    {'.cfg': read_txt, 'category': 'text'},
    {'.com': None, 'category': 'executable'},
    {'.conf': read_txt, 'category': 'text'},
    {'.cpp': read_txt, 'category': 'text'},
    {'.cs': read_txt, 'category': 'text'},
    {'.csv': read_txt, 'category': 'text'},
    {'.db': None, 'category': 'database'},
    {'.dbx': dbx_to_txt, 'category': 'email'},
    {'.doc': doc_to_txt, 'category': 'document'},
    {'.docx': docx_to_txt, 'category': 'document'},
    {'.eml': eml_to_txt, 'category': 'email'},
    {'.enc': image_to_txt, 'category': 'image'},
    {'.epub': None, 'category': 'document'},
    {'.exe': exe_to_txt, 'category': 'executable'},
    {'.htm': html_to_txt, 'category': 'hypertext'},
    {'.html': html_to_txt, 'category': 'hypertext'},
    {'.ini': read_txt, 'category': 'text'},
    {'.ipynb': ipynb_to_txt, 'category': 'notebook'},
    {'.jar': None, 'category': 'executable'},
    {'.java': read_txt, 'category': 'text'},
    {'.jfif': image_to_txt, 'category': 'image'},
    {'.jpeg': image_to_txt, 'category': 'image'},
    {'.jpg': image_to_txt, 'category': 'image'},
    {'.json': read_txt, 'category': 'text'},
    {'.log': read_txt, 'category': 'text'},
    {'.md': read_txt, 'category': 'text'},
    {'.mht': html_to_txt, 'category': 'hypertext'},
    {'.mhtml': html_to_txt, 'category': 'hypertext'},
    {'.mov': audio_video_to_txt, 'category': 'video'},
    {'.mp3': audio_to_txt, 'category': 'audio'},
    {'.mp4': audio_video_to_txt, 'category': 'video'},
    {'.mpeg': audio_video_to_txt, 'category': 'video'},
    {'.msg': msg_to_txt, 'category': 'email'},
    {'.msi': exe_to_txt, 'category': 'executable'},
    {'.ods': None, 'category': 'spreadsheet'},
    {'.odt': None, 'category': 'document'},
    {'.ogg': audio_to_txt, 'category': 'audio'},
    {'.opus': audio_to_txt, 'category': 'audio'},
    {'.pdf': pdf_to_txt, 'category': 'document'},
    {'.png': image_to_txt, 'category': 'image'},
    {'.potx': None, 'category': 'presentation'},
    {'.pps': pptx_to_txt, 'category': 'presentation'},
    {'.pptx': pptx_to_txt, 'category': 'presentation'},
    {'.pst': None, 'category': 'email'},
    {'.py': read_txt, 'category': 'text'},
    {'.rar': None, 'category': 'compressed'},
    {'.rtf': rtf_to_txt, 'category': 'document'},
    {'.sh': read_txt, 'category': 'text'},
    {'.sqlite': None, 'category': 'database'},
    {'.sqlite3': None, 'category': 'database'},
    {'.tsv': read_txt, 'category': 'text'},
    {'.txt': read_txt, 'category': 'text'},
    {'.wav': audio_to_txt, 'category': 'audio'},
    {'.webp': image_to_txt, 'category': 'image'},
    {'.xls': xls_to_txt, 'category': 'spreadsheet'},
    {'.xlsm': xlsx_to_txt, 'category': 'spreadsheet'},
    {'.xlsx': xlsx_to_txt, 'category': 'spreadsheet'},
    {'.xml': read_txt, 'category': 'text'},
    {'.zip': zip_to_txt, 'category': 'compressed'}
]

def run_function(path):
    from os.path import splitext
    for k_type in TYPES:
        for name, func in k_type.items():
            if name == splitext(path)[1].lower():
                value = func(path)
                result = {}
                for k_value in value:
                    if value[k_value] and value[k_value] != '':
                        result[k_value] = value[k_value]
                return result
    return None

@timed
def apply_to_dir(directory):
    import os
    from methods import get_hash
    with open('summary.txt', 'w', encoding=UTF8) as summary_file:
        for entry in os.listdir(directory):
            path = os.path.join(directory, entry)
            if os.path.isfile(path):
                show.error(f'{line_number()} [{path}]...')
                markit = mark(path)
                content = run_function(path)
                if content:
                    if markit and MARKIT in markit:
                        print('=' * 150)
                        print('MARKIT OK')
                        print('=' * 150)
                        if markit[MARKIT] and markit[MARKIT] != '':
                            print('='*150)
                            print('='*150)
                            print('MARKIT MERGE')
                            print('='*150)
                            print('='*150)
                            content = content | markit
                    summary = json.dumps(content)
                    if len(summary.encode(UTF8)) > 8192:
                        summary = json.dumps(txt_to_summary(summary[:8000]))
                else:
                    summary = '{"SUMMARY": "No readable content extracted."}'
                summary_file.write(f'[{get_hash(path, HASH_MD5)}]\t[{path}]\t[{normalize(summary)}]\n')
            elif os.path.isdir(path):
                show.error(f'{line_number()} Directory: {path}')

if __name__ == MAIN:
    import sys
    show.error(f'{line_number()} Analysis started...')

    if len(sys.argv) != 2:
        show.error(f'{line_number()} Usage: python script.py <directory>')
        sys.exit(1)
    apply_to_dir(sys.argv[1])

